# -*- coding: utf-8 -*-
"""Fill a geotechnical workload confirmation sheet from a borehole-column DXF.

Reads a Lizheng-style column DXF plus the borehole inventory, classifies each
layer by the fee-standard complexity table (表3.3—1), and writes:

- a confirmation workbook (.xls) for sections 2–6, copying sections 1, 7 and 8
  from the template sheet unchanged
- a per-hole detail workbook (.xlsx) for checking
"""
import argparse
import json
import re
import sys
from collections import defaultdict

import xlrd
import xlwt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BAND_EDGES = [0, 10, 20, 30, 40, 50, 60]
BAND_NAMES = ["0～10m", "10.1～20m", "20.1～30m", "30.1～40m", "40.1～50m", "50.1～60m"]
CLASS_MARK = {"I": "I级", "II": "Ⅱ级", "III": "Ⅲ级"}
CLASSES = ["I", "II", "III"]
QTY_KEYS = [
    "water_sample", "physical", "atterberg", "compression", "quick_shear",
    "consolidated_quick", "grain", "clay_fraction", "high_pressure", "uu", "cu",
    "salt", "water_analysis", "wave_le15", "wave_mid",
]


def clean(s):
    s = s.replace("\\P", " ").replace("\\A1;", "")
    s = re.sub(r"\\[A-Za-z][^;]*;", "", s)
    return s.replace("{", "").replace("}", "").strip()


def load_dxf(path):
    section = None
    pending = False
    mtexts, texts, inserts = [], [], []
    cur = None
    with open(path, "r", encoding="gbk", errors="replace") as f:
        code = None
        for line in f:
            s = line.rstrip("\r\n")
            if code is None:
                code = s.strip()
                continue
            val = s
            if code == "0" and val == "SECTION":
                pending = True
            elif pending and code == "2":
                section = val
                pending = False
            elif code == "0" and val == "ENDSEC":
                section = None
            if section != "ENTITIES":
                code = None
                continue
            if code == "0":
                if cur:
                    if cur["t"] == "MTEXT" and cur["txt"]:
                        mtexts.append(cur)
                    elif cur["t"] == "TEXT" and cur["txt"]:
                        texts.append(cur)
                    elif cur["t"] == "INSERT":
                        inserts.append(cur)
                cur = {"t": val, "txt": []}
            elif cur is not None:
                if code in ("10", "20"):
                    cur[code] = float(val)
                elif code == "2" and cur["t"] == "INSERT":
                    cur["name"] = val.strip()
                elif code in ("1", "3") and cur["t"] in ("TEXT", "MTEXT"):
                    cur["txt"].append(val)
            code = None
    return mtexts, texts, inserts


def has_word(text, word):
    if word == "密实":
        return "密实" in text.replace("中密", "").replace("稍密", "")
    return word in text


def classify(text):
    """Return (class, reason) using 表3.3—1. Class is I, II or III."""
    raw = text.strip()
    name = re.split(r"[:：]", raw)[0].replace("。", "").strip()
    body = raw
    sand_keys = ("粉细砂", "细中砂", "中粗砂", "粗中砂", "粉砂", "细砂", "中砂", "粗砂", "砾砂")
    is_sand = any(k in name for k in sand_keys) or name.endswith("砂")
    is_clay = any(k in name for k in ("黏土", "粘土", "淤泥"))
    is_silt = "粉土" in name and "黏土" not in name and "粘土" not in name
    is_za = "杂填土" in name
    is_su = "素填土" in name
    is_gravel = any(k in name for k in ("卵石", "碎石", "块石", "漂石"))
    body_clay = ("黏性颗粒" in body) or ("粘性颗粒" in body)
    body_silt = ("粉粒" in body) and not body_clay

    if is_gravel:
        raise ValueError("出现卵石/碎石/块石，确认表只有 I～Ⅲ 级，需人工处理：" + name)
    if is_za:
        if any(k in body for k in ("建筑垃圾", "混凝土", "砖", "灰渣", "块石")):
            return "III", "杂填土，建筑垃圾为主，硬杂质>25%"
        return "I", "杂填土，硬杂质≤10%"
    if is_su:
        if any(k in body for k in ("建筑垃圾", "混凝土", "砖渣", "砖块", "灰渣")):
            if any(k in body for k in ("为主", "大量")):
                return "III", "素填土，硬杂质>25%"
            return "II", "素填土，硬杂质≤25%"
        return "I", "素填土，黏性土及粉土为主，硬杂质≤10%"
    if is_sand:
        return "III", "砂土"
    if is_silt and body_clay:
        if has_word(body, "硬塑") or has_word(body, "坚硬"):
            return "II", "定名为粉土，描述为硬塑黏性土"
        return "I", "定名为粉土，描述为可塑黏性土"
    if is_clay and body_silt:
        if has_word(body, "密实"):
            return "II", "定名为粉质黏土，描述为密实粉土"
        return "I", "定名为粉质黏土，描述为稍密或中密粉土"
    if is_clay or body_clay:
        if has_word(body, "硬塑") or has_word(body, "坚硬"):
            return "II", "硬塑或坚硬黏性土"
        return "I", "可塑、软塑或流塑黏性土"
    if is_silt or body_silt:
        if has_word(body, "密实"):
            return "II", "密实粉土"
        return "I", "稍密或中密粉土"
    raise ValueError("无法归类：" + raw[:80])


def band_cents(top, bot):
    top = round(top, 2)
    bot = round(bot, 2)
    if bot > 60.001:
        raise ValueError("孔深超过 60 m，确认表未设该深度档：%.2f" % bot)
    out = [0] * 6
    for i in range(6):
        a = max(top, BAND_EDGES[i])
        b = min(bot, BAND_EDGES[i + 1])
        if b - a > 0.001:
            out[i] = int(round((b - a) * 100))
    return out


def parse_holes(mtexts, texts):
    ms = [(m["10"], m["20"], clean("".join(m["txt"]))) for m in mtexts]
    ms.sort()
    clusters, cur = [], []
    for item in ms:
        if not cur or item[0] - cur[-1][0] <= 5:
            cur.append(item)
        else:
            clusters.append(cur)
            cur = [item]
    if cur:
        clusters.append(cur)

    num_re = re.compile(r"^-?\d+(\.\d+)?$")
    nums = []
    for t in texts:
        s = "".join(t["txt"]).strip()
        s2 = s[1:] if s.startswith("=") else s
        if num_re.match(s2):
            nums.append((t["10"], t["20"], round(float(s2), 2), s.startswith("=")))

    holes = []
    for cluster in clusters:
        mx = sum(p[0] for p in cluster) / len(cluster)
        cands = [n for n in nums if mx - 45 < n[0] < mx - 10 and not n[3] and 0 < n[2] <= 80]
        byx = defaultdict(list)
        for x, y, v, _eq in cands:
            byx[round(x, 0)].append((y, v))
        best = None
        for pts in byx.values():
            pts_s = sorted(pts, key=lambda z: -z[0])
            vals = [v for _, v in pts_s]
            if len(vals) < 2:
                continue
            if all(vals[i] < vals[i + 1] - 1e-9 for i in range(len(vals) - 1)):
                if best is None or vals[-1] > best[0]:
                    best = (vals[-1], pts_s)
        if not best:
            raise RuntimeError("在 x=%.1f 未找到层底深度列" % mx)
        dys = best[1]
        desc = sorted(cluster, key=lambda z: -z[1])
        used = set()
        paired = []
        for _x, y, t in desc:
            choices = []
            for i, (yd, vd) in enumerate(dys):
                if i in used:
                    continue
                if yd <= y + 2.0:
                    choices.append((y - yd, i, vd))
            if not choices:
                for i, (yd, vd) in enumerate(dys):
                    if i in used:
                        continue
                    choices.append((abs(y - yd), i, vd))
            if not choices:
                continue
            choices.sort()
            _, i, vd = choices[0]
            used.add(i)
            paired.append((round(vd, 2), t))
        paired.sort()
        if len(paired) != len(dys):
            raise RuntimeError("分层描述与层底深度未能逐层配对，x=%.1f" % mx)
        layers = []
        top = 0.0
        for bot, t in paired:
            cls, note = classify(t)
            name = re.split(r"[:：]", t)[0].replace("。", "").strip()
            layers.append({
                "top": top, "bot": bot, "cls": cls, "note": note,
                "name": name, "text": t, "cents": band_cents(top, bot),
            })
            top = bot
        holes.append({"mx": mx, "layers": layers, "depth": layers[-1]["bot"]})
    return holes


def class_of_interval(layers, d0, d1):
    best, best_w, best_ly = None, -1.0, None
    span = 0
    for ly in layers:
        overlap = min(d1, ly["bot"]) - max(d0, ly["top"])
        if overlap > 0.001:
            span += 1
        if overlap > best_w + 1e-9:
            best_w, best, best_ly = overlap, ly["cls"], ly
    if best_ly is None:
        best_ly = layers[-1]
        best = best_ly["cls"]
    return best, best_ly, span > 1


def spt_bin(d0):
    if d0 <= 20:
        return "D≤20"
    if d0 <= 50:
        return "20<D≤50"
    return "D>50"


def load_inventory(path):
    sh = xlrd.open_workbook(path).sheet_by_index(0)
    header_r = None
    for r in range(min(8, sh.nrows)):
        for c in range(sh.ncols):
            if "勘探点编号" in str(sh.cell_value(r, c)).replace("\n", ""):
                header_r = r
                break
        if header_r is not None:
            break
    if header_r is None:
        raise RuntimeError("勘探点一览表中未找到“勘探点编号”")
    mapping = {}
    for rr in (header_r, header_r + 1):
        if rr >= sh.nrows:
            continue
        for c in range(sh.ncols):
            t = str(sh.cell_value(rr, c)).replace("\n", "").replace(" ", "")
            if "勘探点编号" in t:
                mapping["name"] = c
            elif "类型" in t:
                mapping["kind"] = c
            elif "钻探" in t:
                mapping["depth"] = c
            elif "原" in t and "样" in t:
                mapping["uy"] = c
            elif "扰" in t and "样" in t:
                mapping["dy"] = c
            elif "标贯" in t:
                mapping["spt"] = c
    missing = [k for k in ("name", "kind", "depth", "uy", "dy", "spt") if k not in mapping]
    if missing:
        raise RuntimeError("勘探点一览表缺少列：" + "、".join(missing))

    def num(r, key):
        v = sh.cell_value(r, mapping[key])
        return float(v) if v not in ("", None) else 0.0

    rows = []
    for r in range(header_r + 2, sh.nrows):
        name = str(sh.cell_value(r, mapping["name"])).strip()
        if not name or name == "合计":
            continue
        rows.append({
            "name": name,
            "kind": str(sh.cell_value(r, mapping["kind"]) or "").strip(),
            "depth": num(r, "depth"),
            "uy": num(r, "uy"),
            "dy": num(r, "dy"),
            "spt": num(r, "spt"),
        })
    return rows


def collect_tests(holes, texts, inserts):
    rng_re = re.compile(r"^(\d+\.\d+)\s*[-~～]\s*(\d+\.\d+)$")
    ranges = []
    for t in texts:
        s = "".join(t["txt"]).strip().replace(" ", "")
        m = rng_re.match(s)
        if not m:
            continue
        a, b = round(float(m.group(1)), 2), round(float(m.group(2)), 2)
        ranges.append((t["10"], t["20"], a, b, round(b - a, 2)))
    anchors = [h["mx"] for h in holes]

    def hole_index(x):
        j = min(range(len(anchors)), key=lambda k: abs(anchors[k] - x))
        if abs(anchors[j] - x) > 180:
            raise RuntimeError("有深度标注无法归入任一钻孔，x=%.1f" % x)
        return j

    spt_rows = []
    for x, y, a, b, ln in ranges:
        if abs(ln - 0.3) > 0.02:
            continue
        j = hole_index(x)
        cls, ly, cross = class_of_interval(holes[j]["layers"], a, b)
        spt_rows.append({
            "j": j, "a": a, "b": b, "cls": cls, "bin": spt_bin(a),
            "name": ly["name"], "top": ly["top"], "bot": ly["bot"], "cross": cross,
        })

    qy1 = [(i["10"], i["20"]) for i in inserts if i.get("name") == "gkqy1.dwgZZT"]
    by_ranges, by_qy = defaultdict(list), defaultdict(list)
    for x, y, a, b, ln in ranges:
        if abs(ln - 0.2) < 0.02 or abs(ln - 0.1) < 0.02:
            by_ranges[hole_index(x)].append((x, y, a, b, ln))
    for x, y in qy1:
        by_qy[hole_index(x)].append((x, y))

    sample_rows = []
    for j in sorted(set(by_ranges) | set(by_qy)):
        pool = by_ranges[j][:]
        used = set()
        for x, y in sorted(by_qy.get(j, []), key=lambda p: -p[1]):
            best_i, best_dy = None, 1e18
            for i, (_rx, ry, _a, _b, _ln) in enumerate(pool):
                if i in used:
                    continue
                dy = abs(y - ry)
                if dy < best_dy:
                    best_dy, best_i = dy, i
            if best_i is None or best_dy > 8:
                raise RuntimeError("原状土样符号无法对应取样深度，孔序号 %d" % (j + 1))
            used.add(best_i)
            a, b = pool[best_i][2], pool[best_i][3]
            cls, ly, cross = class_of_interval(holes[j]["layers"], a, b)
            sample_rows.append({
                "j": j, "kind": "原状", "a": a, "b": b,
                "band": "D≤30" if a <= 30 else "D＞30",
                "cls": cls, "name": ly["name"], "cross": cross,
            })
        for i, (_rx, _ry, a, b, _ln) in enumerate(pool):
            if i in used:
                continue
            cls, ly, cross = class_of_interval(holes[j]["layers"], a, b)
            sample_rows.append({
                "j": j, "kind": "扰动", "a": a, "b": b, "band": "",
                "cls": cls, "name": ly["name"], "cross": cross,
            })
    return spt_rows, sample_rows


def check_alignment(holes, inv, spt_rows, sample_rows):
    errors = []
    if len(holes) != len(inv):
        errors.append("柱状图孔数 %d 与一览表 %d 不一致" % (len(holes), len(inv)))
        return errors
    spt_n = defaultdict(int)
    for row in spt_rows:
        spt_n[row["j"]] += 1
    uy_n, dy_n = defaultdict(int), defaultdict(int)
    for row in sample_rows:
        if row["kind"] == "原状":
            uy_n[row["j"]] += 1
        else:
            dy_n[row["j"]] += 1
    for i, (h, info) in enumerate(zip(holes, inv)):
        if abs(h["depth"] - info["depth"]) > 0.02:
            errors.append("%s 孔深 图面 %.2f / 一览表 %.2f" % (info["name"], h["depth"], info["depth"]))
        if spt_n[i] != int(info["spt"]):
            errors.append("%s 标贯 图面 %d / 一览表 %d" % (info["name"], spt_n[i], int(info["spt"])))
        if uy_n[i] != int(info["uy"]):
            errors.append("%s 原状 图面 %d / 一览表 %d" % (info["name"], uy_n[i], int(info["uy"])))
        if dy_n[i] != int(info["dy"]):
            errors.append("%s 扰动 图面 %d / 一览表 %d" % (info["name"], dy_n[i], int(info["dy"])))
    if len(errors) > 12:
        errors = errors[:12] + ["……其余 %d 条未列出" % (len(errors) - 12)]
    return errors


def drill_vector(holes):
    cents = [0] * 18
    for h in holes:
        for ly in h["layers"]:
            for bi, c in enumerate(ly["cents"]):
                cents[bi * 3 + CLASSES.index(ly["cls"])] += c
    return [c / 100.0 for c in cents]


def spt_vector(spt_rows):
    order = []
    for depth_key in ("D≤20", "20<D≤50", "D>50"):
        for cls in CLASSES:
            order.append(sum(1 for r in spt_rows if r["bin"] == depth_key and r["cls"] == cls))
    return order


def section_no(value):
    try:
        f = float(value)
    except (TypeError, ValueError):
        return None
    if f in (1, 2, 3, 4, 5, 6, 7, 8):
        return str(int(f))
    return None


def read_kept_sections(template, sheet_name):
    wb = xlrd.open_workbook(template)
    if sheet_name not in wb.sheet_names():
        raise RuntimeError("模板中没有工作表“%s”，现有：%s" % (sheet_name, "、".join(wb.sheet_names())))
    sh = wb.sheet_by_name(sheet_name)
    blocks = {}
    current = None
    buf = []
    for r in range(sh.nrows):
        no = sh.cell_value(r, 0) if sh.ncols > 0 else ""
        name = sh.cell_value(r, 1) if sh.ncols > 1 else ""
        unit = sh.cell_value(r, 2) if sh.ncols > 2 else ""
        qty = sh.cell_value(r, 3) if sh.ncols > 3 else ""
        if "申报" in str(no) or "申报" in str(name):
            if current:
                blocks[current] = buf
            break
        sec = section_no(no)
        if sec:
            if current:
                blocks[current] = buf
            current = sec
            buf = [(no, name, unit, qty)]
        elif current is not None and (str(no).strip() or str(name).strip()):
            buf.append((no, name, unit, qty))
    if current and current not in blocks:
        blocks[current] = buf
    for key in ("1", "7", "8"):
        if key not in blocks:
            raise RuntimeError("模板工作表“%s”缺少第 %s 项，不能原样保留" % (sheet_name, key))
    return blocks


def circled(i):
    return chr(0x2460 + i)


def build_form_rows(kept, drill, spt, uy_le, uy_gt, qty):
    rows = []

    def add_section(block):
        for i, (no, name, unit, q) in enumerate(block):
            if i == 0:
                rows.append(("section", section_label(no), str(name).strip(), None, None))
            else:
                rows.append(("item", str(no).strip(), str(name).strip(), str(unit).strip(), float(q or 0)))

    def section_label(no):
        try:
            return "%.1f" % float(no)
        except (TypeError, ValueError):
            return str(no).strip()

    add_section(kept["1"])
    rows.append(("section", "2.0", "岩土工程勘探", None, None))
    n = 0
    for band in BAND_NAMES:
        for cls in CLASSES:
            rows.append(("item", circled(n), "%s（%s）" % (band, CLASS_MARK[cls]), "m", drill[n]))
            n += 1
    rows.append(("total", "总钻探量（m）", None, None, round(sum(drill), 1)))
    rows.append(("section", "3.0", "原位测试", None, None))
    spt_labels = []
    for depth, pretty in (("D≤20", "D≤20"), ("20<D≤50", "20＜D≤50"), ("D>50", "D＞50")):
        for cls, mark in (("I", "I类"), ("II", "Ⅱ类"), ("III", "Ⅲ类")):
            spt_labels.append("标准贯入试验（%s%s）" % (mark, pretty))
    for i, lab in enumerate(spt_labels):
        rows.append(("item", circled(i), lab, "次", float(spt[i])))
    rows.append(("section", "4.0", "取土、水、岩", None, None))
    for i, (lab, unit, q) in enumerate((
        ("原状土样D≤30", "组", float(uy_le)),
        ("原状土样D＞30", "组", float(uy_gt)),
        ("扰动土样", "组", float(qty.get("disturbed_used"))),
        ("水样", "组", float(qty["water_sample"])),
    )):
        rows.append(("item", circled(i), lab, unit, q))
    rows.append(("section", "5.0", "岩土试验", None, None))
    tests = [
        ("含水量、容量、   比重", qty["physical"]),
        ("液塑限", qty["atterberg"]),
        ("压缩", qty["compression"]),
        ("快剪", qty["quick_shear"]),
        ("固结快剪", qty["consolidated_quick"]),
        ("颗分", qty["grain"]),
        ("颗分（黏粒含量）", qty["clay_fraction"]),
        ("高压固结试验", qty["high_pressure"]),
        ("静三轴剪切（UU）", qty["uu"]),
        ("静三轴剪切（CU）", qty["cu"]),
        ("易溶盐分析", qty["salt"]),
        ("水简分析", qty["water_analysis"]),
    ]
    for i, (lab, q) in enumerate(tests):
        rows.append(("item", circled(i), lab, "组", float(q)))
    rows.append(("section", "6.0", "波速测试", None, None))
    rows.append(("item", circled(0), "单孔法D≤15m", "m", float(qty["wave_le15"])))
    rows.append(("item", circled(1), "单孔法15<D≤30m", "m", float(qty["wave_mid"])))
    add_section(kept["7"])
    add_section(kept["8"])
    return rows


def write_form_xls(path, title, sheet_name, form_rows):
    wb = xlwt.Workbook()
    ws = wb.add_sheet(sheet_name[:31])
    font = "宋体"
    title_s = xlwt.easyxf(
        "font: name %s, height 280, bold on; align: horiz center, vert center, wrap on;" % font
    )
    header = xlwt.easyxf(
        "font: name %s, height 220, bold on; align: horiz center, vert center, wrap on;"
        "borders: left thin, right thin, top thin, bottom thin;" % font
    )
    sec = xlwt.easyxf(
        "font: name %s, height 220, bold on; align: horiz center, vert center;"
        "borders: left thin, right thin, top thin, bottom thin;" % font
    )
    sec_left = xlwt.easyxf(
        "font: name %s, height 220, bold on; align: horiz left, vert center;"
        "borders: left thin, right thin, top thin, bottom thin;" % font
    )
    cell = xlwt.easyxf(
        "font: name %s, height 220; align: horiz center, vert center, wrap on;"
        "borders: left thin, right thin, top thin, bottom thin;" % font
    )
    left = xlwt.easyxf(
        "font: name %s, height 220; align: horiz left, vert center, wrap on;"
        "borders: left thin, right thin, top thin, bottom thin;" % font
    )
    num = xlwt.easyxf(
        "font: name %s, height 220; align: horiz center, vert center;"
        "borders: left thin, right thin, top thin, bottom thin;" % font,
        num_format_str="0.0",
    )
    foot = xlwt.easyxf("font: name %s, height 220; align: horiz left, vert center, wrap on;" % font)
    for i, w in enumerate((4200, 9800, 4200, 4200)):
        ws.col(i).width = w
    ws.write_merge(0, 0, 0, 3, title, title_s)
    ws.row(0).height_mismatch = True
    ws.row(0).height = 1080
    for c, text in enumerate(["编号", "工程及费用项目", "计算单位", "工程数量"]):
        ws.write_merge(1, 2, c, c, text, header)
    ws.row(1).height_mismatch = True
    ws.row(1).height = 400
    ws.row(2).height_mismatch = True
    ws.row(2).height = 400
    r = 3
    for kind, no, name, unit, qty in form_rows:
        ws.row(r).height_mismatch = True
        ws.row(r).height = 740
        if kind == "total":
            ws.write_merge(r, r, 0, 2, no, sec_left)
            ws.write(r, 3, qty, num)
        elif kind == "section":
            ws.write(r, 0, no, sec)
            ws.write_merge(r, r, 1, 3, name, sec_left)
        else:
            ws.write(r, 0, no, cell)
            ws.write(r, 1, name, left)
            ws.write(r, 2, unit, cell)
            ws.write(r, 3, qty, num)
        r += 1
    ws.row(r).height_mismatch = True
    ws.row(r).height = 1400
    ws.write_merge(
        r, r, 0, 3,
        "\n           申报人：                                         确认人：\n\n"
        "           日  期：                                         日  期：\n",
        foot,
    )
    wb.save(path)


def write_detail(path, holes, inv, spt_rows, sample_rows, drill, spt, uy_le, uy_gt, qty):
    wb = Workbook()
    thin = Border(
        left=Side(style="thin", color="B0B0B0"), right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"), bottom=Side(style="thin", color="B0B0B0"),
    )
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(name="宋体", size=10, bold=True, color="FFFFFF")
    cell_font = Font(name="宋体", size=10)
    total_font = Font(name="宋体", size=10, bold=True)
    total_fill = PatternFill("solid", fgColor="D6E3F0")
    fills = {
        "I": PatternFill("solid", fgColor="E2EFDA"),
        "II": PatternFill("solid", fgColor="FFF2CC"),
        "III": PatternFill("solid", fgColor="FCE4D6"),
    }
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def style_header(ws, ncol, height=32):
        ws.row_dimensions[1].height = height
        ws.freeze_panes = "A2"
        ws.print_title_rows = "1:1"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.horizontalCentered = True
        ws.sheet_view.showGridLines = False
        ws.page_margins.left = 0.4
        ws.page_margins.right = 0.4
        ws.page_margins.top = 0.6
        ws.page_margins.bottom = 0.5
        ws.oddFooter.right.text = "第 &P 页 / 共 &N 页"
        for c in range(1, ncol + 1):
            cell = ws.cell(1, c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = center
            cell.border = thin

    def put(ws, r, c, value, font=None, fill=None, align=None, fmt=None):
        cell = ws.cell(r, c, value)
        cell.font = font or cell_font
        cell.border = thin
        cell.alignment = align or center
        if fill is not None:
            cell.fill = fill
        if fmt:
            cell.number_format = fmt
        return cell

    ws = wb.active
    ws.title = "与确认表核对"
    for c, h in enumerate(["项目", "明细合计", "确认表", "差值"], 1):
        ws.cell(1, c, h)
    style_header(ws, 4, 22)
    labels, detail, expect = [], [], []
    for band in BAND_NAMES:
        for mark in ("I级", "Ⅱ级", "Ⅲ级"):
            labels.append("钻探 %s（%s） m" % (band, mark))
    detail += drill
    expect += drill
    for lab, n in zip(
        ["标贯 I类 D≤20", "标贯 Ⅱ类 D≤20", "标贯 Ⅲ类 D≤20",
         "标贯 I类 20＜D≤50", "标贯 Ⅱ类 20＜D≤50", "标贯 Ⅲ类 20＜D≤50",
         "标贯 I类 D＞50", "标贯 Ⅱ类 D＞50", "标贯 Ⅲ类 D＞50"],
        spt,
    ):
        labels.append(lab + " 次")
        detail.append(n)
        expect.append(n)
    tail = [
        ("原状土样 D≤30 组", uy_le, uy_le),
        ("原状土样 D＞30 组", uy_gt, uy_gt),
        ("扰动土样 组", qty["disturbed_used"], qty["disturbed_used"]),
        ("水样 组", qty["water_sample"], qty["water_sample"]),
        ("含水量、容量、比重 组", qty["physical"], qty["physical"]),
        ("液塑限 组", qty["atterberg"], qty["atterberg"]),
        ("压缩 组", qty["compression"], qty["compression"]),
        ("快剪 组", qty["quick_shear"], qty["quick_shear"]),
        ("固结快剪 组", qty["consolidated_quick"], qty["consolidated_quick"]),
        ("颗分 组", qty["grain"], qty["grain"]),
        ("颗分（黏粒含量）组", qty["clay_fraction"], qty["clay_fraction"]),
        ("高压固结 组", qty["high_pressure"], qty["high_pressure"]),
        ("三轴UU 组", qty["uu"], qty["uu"]),
        ("三轴CU 组", qty["cu"], qty["cu"]),
        ("易溶盐 组", qty["salt"], qty["salt"]),
        ("水简分析 组", qty["water_analysis"], qty["water_analysis"]),
        ("单孔波速 D≤15m", qty["wave_le15"], qty["wave_le15"]),
        ("单孔波速 15＜D≤30m", qty["wave_mid"], qty["wave_mid"]),
    ]
    for lab, a, b in tail:
        labels.append(lab)
        detail.append(a)
        expect.append(b)
    for i, (lab, a, b) in enumerate(zip(labels, detail, expect), 2):
        put(ws, i, 1, lab, align=left)
        put(ws, i, 2, a, fmt="0.0")
        put(ws, i, 3, b, fmt="0.0")
        diff = round(float(a) - float(b), 2)
        put(ws, i, 4, diff, fmt="0.0")
    for col, width in zip("ABCD", (36, 14, 14, 12)):
        ws.column_dimensions[col].width = width
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToHeight = 1
    ws.auto_filter.ref = "A1:D%d" % (1 + len(labels))

    # 钻孔汇总
    ws = wb.create_sheet("钻孔汇总")
    headers = ["序号", "孔号", "孔类型", "一览表孔深", "图面孔深"]
    for band in BAND_NAMES:
        for mark in ("I", "Ⅱ", "Ⅲ"):
            headers.append("%s\n%s级" % (band, mark))
    headers += [
        "进尺合计", "一览表标贯", "标贯合计",
        "标贯I\nD≤20", "标贯Ⅱ\nD≤20", "标贯Ⅲ\nD≤20",
        "标贯I\n20<D≤50", "标贯Ⅱ\n20<D≤50", "标贯Ⅲ\n20<D≤50",
        "标贯I\nD>50", "标贯Ⅱ\nD>50", "标贯Ⅲ\nD>50",
        "一览表原状", "原状合计", "原状D≤30", "原状D＞30", "一览表扰动", "扰动",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    style_header(ws, len(headers), 36)
    spt_by = defaultdict(lambda: [0] * 9)
    bins = ["D≤20", "20<D≤50", "D>50"]
    for row in spt_rows:
        idx = bins.index(row["bin"]) * 3 + CLASSES.index(row["cls"])
        spt_by[row["j"]][idx] += 1
    uy_by = defaultdict(lambda: [0, 0])
    dy_by = defaultdict(int)
    for row in sample_rows:
        if row["kind"] == "原状":
            uy_by[row["j"]][0 if row["band"] == "D≤30" else 1] += 1
        else:
            dy_by[row["j"]] += 1
    for i, (h, info) in enumerate(zip(holes, inv)):
        r = i + 2
        cents = [0] * 18
        for ly in h["layers"]:
            for bi, c in enumerate(ly["cents"]):
                cents[bi * 3 + CLASSES.index(ly["cls"])] += c
        sc = spt_by[i]
        uy = uy_by[i]
        vals = [i + 1, info["name"], info["kind"], info["depth"], h["depth"]]
        vals += [c / 100 if c else None for c in cents]
        vals += [h["depth"], info["spt"] or None, sum(sc) or None]
        vals += [n or None for n in sc]
        vals += [info["uy"] or None, (uy[0] + uy[1]) or None, uy[0] or None, uy[1] or None,
                 info["dy"] or None, dy_by[i] or None]
        for c, v in enumerate(vals, 1):
            fmt = "0.0" if isinstance(v, float) else None
            put(ws, r, c, v, fmt=fmt)
        ws.row_dimensions[r].height = 18
    tr = len(holes) + 2
    for c in range(1, len(headers) + 1):
        put(ws, tr, c, None, font=total_font, fill=total_fill)
    put(ws, tr, 2, "合计", font=total_font, fill=total_fill, align=left)
    put(ws, tr, 4, round(sum(x["depth"] for x in inv), 1), font=total_font, fill=total_fill, fmt="0.0")
    put(ws, tr, 5, round(sum(h["depth"] for h in holes), 1), font=total_font, fill=total_fill, fmt="0.0")
    for k, m in enumerate(drill):
        put(ws, tr, 6 + k, m, font=total_font, fill=total_fill, fmt="0.0")
    base = 24
    put(ws, tr, base, round(sum(drill), 1), font=total_font, fill=total_fill, fmt="0.0")
    put(ws, tr, base + 1, int(sum(x["spt"] for x in inv)), font=total_font, fill=total_fill)
    put(ws, tr, base + 2, sum(spt), font=total_font, fill=total_fill)
    for k, n in enumerate(spt):
        put(ws, tr, base + 3 + k, n, font=total_font, fill=total_fill)
    put(ws, tr, base + 12, int(sum(x["uy"] for x in inv)), font=total_font, fill=total_fill)
    put(ws, tr, base + 13, uy_le + uy_gt, font=total_font, fill=total_fill)
    put(ws, tr, base + 14, uy_le, font=total_font, fill=total_fill)
    put(ws, tr, base + 15, uy_gt, font=total_font, fill=total_fill)
    put(ws, tr, base + 16, int(sum(x["dy"] for x in inv)), font=total_font, fill=total_fill)
    put(ws, tr, base + 17, sum(1 for r in sample_rows if r["kind"] == "扰动"), font=total_font, fill=total_fill)
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(holes) + 1)
    widths = [6, 16, 18, 12, 12] + [9] * 18 + [10, 12, 10] + [9] * 9 + [12, 10, 11, 11, 12, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.page_setup.fitToWidth = 1
    ws.oddHeader.center.text = "钻孔汇总"

    ws = wb.create_sheet("分层明细")
    headers = ["序号", "孔号", "孔类型", "层序", "层顶(m)", "层底(m)", "层厚(m)", "定名", "复杂程度", "判定依据",
               "0～10m", "10.1～20m", "20.1～30m", "30.1～40m", "40.1～50m", "50.1～60m", "地层描述"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    style_header(ws, len(headers), 22)
    r = 2
    for i, (h, info) in enumerate(zip(holes, inv)):
        for k, ly in enumerate(h["layers"], 1):
            vals = [i + 1, info["name"], info["kind"], k, ly["top"], ly["bot"],
                    round(ly["bot"] - ly["top"], 2), ly["name"], ly["cls"] + "级", ly["note"]]
            vals += [c / 100 if c else None for c in ly["cents"]]
            vals.append(ly["text"])
            for c, v in enumerate(vals, 1):
                align = left if c in (10, 17) else center
                fmt = "0.00" if isinstance(v, float) else None
                put(ws, r, c, v, align=align, fmt=fmt, fill=fills.get(ly["cls"]) if c == 9 else None)
            ws.row_dimensions[r].height = 18
            r += 1
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), r - 1)
    for i, w in enumerate([6, 16, 16, 6, 10, 10, 10, 12, 10, 32, 10, 12, 12, 12, 12, 12, 62], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.oddHeader.center.text = "分层明细"

    ws = wb.create_sheet("标贯明细")
    headers = ["序号", "孔号", "孔类型", "起始深度(m)", "终止深度(m)", "深度档", "复杂程度", "所在层定名", "层顶(m)", "层底(m)", "是否跨层"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    style_header(ws, len(headers), 22)
    for n, row in enumerate(sorted(spt_rows, key=lambda z: (z["j"], z["a"])), 1):
        info = inv[row["j"]]
        vals = [row["j"] + 1, info["name"], info["kind"], row["a"], row["b"], row["bin"],
                row["cls"] + "级", row["name"], row["top"], row["bot"], "跨层" if row["cross"] else None]
        for c, v in enumerate(vals, 1):
            fmt = "0.00" if isinstance(v, float) else None
            put(ws, n + 1, c, v, fmt=fmt, fill=fills.get(row["cls"]) if c == 7 else None)
        ws.row_dimensions[n + 1].height = 18
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(spt_rows) + 1)
    for i, w in enumerate([6, 16, 16, 14, 14, 14, 10, 14, 10, 10, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.oddHeader.center.text = "标贯明细"

    ws = wb.create_sheet("取样明细")
    headers = ["序号", "孔号", "孔类型", "取样类型", "起始深度(m)", "终止深度(m)", "深度档", "所在层定名", "复杂程度", "是否跨层"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    style_header(ws, len(headers), 22)
    ordered = sorted(sample_rows, key=lambda z: (z["j"], 0 if z["kind"] == "原状" else 1, z["a"]))
    for n, row in enumerate(ordered, 1):
        info = inv[row["j"]]
        vals = [row["j"] + 1, info["name"], info["kind"], row["kind"], row["a"], row["b"],
                row["band"] or None, row["name"], row["cls"] + "级", "跨层" if row["cross"] else None]
        for c, v in enumerate(vals, 1):
            fmt = "0.00" if isinstance(v, float) else None
            put(ws, n + 1, c, v, fmt=fmt, fill=fills.get(row["cls"]) if c == 9 else None)
        ws.row_dimensions[n + 1].height = 18
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(ordered) + 1)
    for i, w in enumerate([6, 16, 16, 10, 14, 14, 12, 14, 10, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.oddHeader.center.text = "取样明细"
    wb.save(path)


def load_qty(path, disturbed_chart):
    qty = json.load(open(path, encoding="utf-8"))
    missing = [k for k in QTY_KEYS if k not in qty]
    if missing:
        raise RuntimeError("数量 JSON 缺少：" + "、".join(missing))
    if "disturbed" in qty and int(qty["disturbed"]) != disturbed_chart:
        raise RuntimeError("一览表扰动土样 %s 与柱状图 %d 不一致" % (qty["disturbed"], disturbed_chart))
    qty["disturbed_used"] = disturbed_chart
    return qty


def main():
    p = argparse.ArgumentParser(description="由钻孔柱状图填写勘察工作量确认表，并输出逐孔明细")
    p.add_argument("--dxf", required=True)
    p.add_argument("--inventory", required=True)
    p.add_argument("--template", required=True)
    p.add_argument("--keep-sheet", required=True, help="从该表原样复制第1、7、8项")
    p.add_argument("--title", required=True)
    p.add_argument("--sheet-name", required=True, help="确认表工作表名，不超过 31 字")
    p.add_argument("--qty", required=True, help="第4项水样以及第5、6项的 JSON")
    p.add_argument("--out-form", required=True)
    p.add_argument("--out-detail", required=True)
    args = p.parse_args()

    mtexts, texts, inserts = load_dxf(args.dxf)
    holes = parse_holes(mtexts, texts)
    inv = load_inventory(args.inventory)
    spt_rows, sample_rows = collect_tests(holes, texts, inserts)
    errors = check_alignment(holes, inv, spt_rows, sample_rows)
    if errors:
        sys.stderr.write("柱状图与勘探点一览表未对齐，未写出文件：\n")
        sys.stderr.write("\n".join(errors) + "\n")
        return 1
    drill = drill_vector(holes)
    spt = spt_vector(spt_rows)
    band0 = sum(min(h["depth"], 10) for h in holes)
    if abs(band0 - sum(drill[0:3])) > 0.02:
        sys.stderr.write("0～10 m 进尺 %.2f 与各孔 min(孔深,10) 之和 %.2f 不一致\n" % (sum(drill[0:3]), band0))
        return 1
    if abs(sum(drill) - sum(h["depth"] for h in holes)) > 0.02:
        sys.stderr.write("分档进尺合计与孔深合计不一致\n")
        return 1
    uy_le = sum(1 for r in sample_rows if r["kind"] == "原状" and r["band"] == "D≤30")
    uy_gt = sum(1 for r in sample_rows if r["kind"] == "原状" and r["band"] == "D＞30")
    dy_n = sum(1 for r in sample_rows if r["kind"] == "扰动")
    qty = load_qty(args.qty, dy_n)
    kept = read_kept_sections(args.template, args.keep_sheet)
    form_rows = build_form_rows(kept, drill, spt, uy_le, uy_gt, qty)
    write_form_xls(args.out_form, args.title, args.sheet_name, form_rows)
    write_detail(args.out_detail, holes, inv, spt_rows, sample_rows, drill, spt, uy_le, uy_gt, qty)
    swapped = sum(1 for h in holes for ly in h["layers"] if ly["note"].startswith("定名"))
    cross = sum(1 for r in spt_rows if r["cross"])
    print(json.dumps({
        "holes": len(holes),
        "depth_m": round(sum(h["depth"] for h in holes), 2),
        "drill": [round(x, 1) for x in drill],
        "spt": spt,
        "spt_total": sum(spt),
        "undisturbed_le30": uy_le,
        "undisturbed_gt30": uy_gt,
        "disturbed": dy_n,
        "swapped_description_layers": swapped,
        "spt_cross_layer": cross,
        "kept_sheet": args.keep_sheet,
        "form": args.out_form,
        "detail": args.out_detail,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
